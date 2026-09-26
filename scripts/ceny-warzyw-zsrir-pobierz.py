#!/usr/bin/env python3
"""Odświeża ceny WARZYW w database/data/ceny_skladnikow.csv z MRiRW/ZSRIR (D-286, część 3).

DLACZEGO NIE GUS. Bank Danych Lokalnych GUS (temat P1466, źródło pozostałych
wierszy cennika) NIE PODAJE dziś cen warzyw — seria miesięczna z cenami
ziemniaków, cebuli i marchwi kończy się w 2019 roku. Zamiennik: Zintegrowany
System Rolniczej Informacji Rynkowej (ZSRIR) Ministerstwa Rolnictwa i Rozwoju
Wsi, publikowany jako otwarte dane na dane.gov.pl (zbiór 912, licencja CC BY
4.0 / domena publiczna, format xlsx, cotygodniowa aktualizacja).

URUCHAMIA SIĘ AUTOMATYCZNIE, RAZ NA TYDZIEŃ, W GITHUB ACTIONS
(`.github/workflows/ceny-warzyw-auto.yml`, decyzja właściciela z 26.09.2026,
D-286 część 3) — oraz ręcznie, na komputerze osoby prowadzącej, dokładnie
jak `ceny-gus-pobierz.py`. W obu przypadkach skrypt zmienia TYLKO plik
w repozytorium i TYLKO w gałęzi/PR-ze do przeglądu — nigdy prosto na `main`
i nigdy w samej produkcji. Po scaleniu PR-a produkcja wczytuje nowy plik
komendą `php artisan kuking:ceny-skladnikow`. Produkcja SAMA niczego
z sieci nie pobiera — cotygodniowy rytm dotyczy wyłącznie tego skryptu
w CI, nie aplikacji.

CO SKRYPT AKTUALIZUJE — DWA ARKUSZE TEGO SAMEGO BIULETYNU.

1. „ZAKUP WARZ DETAL - DO 2 KG” (`MAPA` niżej): średnia KRAJOWA cena zakupu
   warzyw PRZEZ PODMIOTY HANDLU DETALICZNEGO w opakowaniach do 2 kg —
   najbliższy oficjalny, cotygodniowy odpowiednik ceny detalicznej, jaki
   ZSRIR ma. Dotyczy pięciu wierszy: `ziemniaki`, `cebula`, `marchew`,
   `papryka_czerwona`, `pomidor`.

2. „HURT WARZ” (`MAPA_HURT` niżej, decyzja właściciela z 26.09.2026, D-286
   część 3, „Uzupełnienie"): siedem warzyw, których arkusz detaliczny NIE
   notuje — `kapusta`, `buraki`, `por`, `seler`, `pietruszka`, `salata`,
   `ogorek`. Ich cena to ŚREDNIA z min–max każdego z pięciu rynków hurtowych
   (Bronisze, Kalisz, Łódź, Poznań, Rzeszów), a potem średnia arytmetyczna
   tych pięciu wartości, pomnożona przez jeden nazwany, udokumentowany
   przelicznik `App\Domain\Recipes\Koszt\SzacunekKosztuZCen::MNOZNIK_HURT_DETAL`.
   Ten skrypt NIE trzyma własnej kopii liczby 1,6 — czyta ją wprost z pliku
   PHP (`mnoznik_hurt_detal`), więc zmiana stałej w jednym miejscu zmienia
   też to, co ten skrypt policzy, a test zgodności pilnuje, żeby ktoś nie
   zapomniał, że oba języki muszą się zgadzać.

Oba arkusze schodzą w JEDNYM przebiegu i JEDNYM PR-ze — nie ma powodu, żeby
recenzent widział dwie osobne zmiany tego samego cennika w tym samym
tygodniu (D-286 część 3, decyzja właściciela z 26.09.2026: „Jeden PR
tygodniowo ma obejmować wszystkie 12 warzyw”).

Nie dodaje nowych wierszy: nowy klucz (kolejne warzywo) dopisuje się ręcznie
w PR-ze razem z `wzorce`/`wyklucz`/miarami domowymi — to decyzje o dopasowaniu
tekstu składnika, których żaden skrypt nie powinien zgadywać.

Użycie:
    python3 scripts/ceny-warzyw-zsrir-pobierz.py             # podgląd różnic
    python3 scripts/ceny-warzyw-zsrir-pobierz.py --zapisz    # zapis do pliku

Wymaga: pip install openpyxl (w CI instaluje to workflow; ręcznie na
komputerze osoby prowadzącej — nie jest to zależność aplikacji).

Uwaga o rytmie: ceny z tego skryptu odświeżają się co tydzień — tyle wynosi
rytm biuletynu ZSRIR, a ceny warzyw sezonowych (pomidor, papryka) potrafią
się w ciągu paru tygodni mocno zmienić. Cennik mięsa i nabiału (GUS)
odświeża się nadal raz na kwartał — ten inny rytm jest świadomy, nie
pomyłką: GUS publikuje ŚREDNIE ROCZNE, więc częstszy odczyt niczego by
nie zmienił.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import urllib.request
from pathlib import Path

PLIK = Path(__file__).resolve().parent.parent / "database" / "data" / "ceny_skladnikow.csv"
PLIK_MNOZNIKA = (
    Path(__file__).resolve().parent.parent / "app" / "Domain" / "Recipes" / "Koszt" / "SzacunekKosztuZCen.php"
)
API_ZASOBY = "https://api.dane.gov.pl/1.4/datasets/912/resources?sort=-data_date&page=1&per_page=10"

ARKUSZ = "ZAKUP WARZ DETAL - DO 2 KG"
SEKCJA = "WARZYWA krajowe - opakowania do 2 kg"

# klucz w ceny_skladnikow.csv → dokładna etykieta w kolumnie „TOWAR” arkusza.
MAPA = {
    "ziemniaki": "Ziemniaki",
    "cebula": "Cebula biała",
    "marchew": "Marchew",
    "papryka_czerwona": "Papryka czerwona",
    "pomidor": "Pomidory okrągłe",
}

ARKUSZ_HURT = "HURT WARZ"
SEKCJA_HURT = "WARZYWA krajowe - hurt"

# Pięć rynków hurtowych z biuletynu ZSRIR — kolejność musi zgadzać się
# z kolejnością par (min, max) w kolumnach arkusza `ARKUSZ_HURT`.
RYNKI_HURT = ["Bronisze", "Kalisz", "Łódź", "Poznań", "Rzeszów"]

# klucz w ceny_skladnikow.csv → (etykieta w kolumnie „TOWAR” arkusza HURT WARZ, jednostka).
# `szt` = arkusz notuje ten towar za sztukę, nie za kg (por, sałata) — D-286 część 3.
MAPA_HURT: dict[str, tuple[str, str]] = {
    "kapusta": ("Kapusta biala", "kg"),
    "buraki": ("Buraki cwiklowe", "kg"),
    "por": ("Por", "szt"),
    "seler": ("Seler korzeniowy", "kg"),
    "pietruszka": ("Pietruszka (korzeń)", "kg"),
    "salata": ("Sałata (głowa)", "szt"),
    "ogorek": ("Ogórek gruntowy", "kg"),
}


def mnoznik_hurt_detal() -> float:
    """Czyta `MNOZNIK_HURT_DETAL` z JEDYNEGO źródła prawdy — stałej w PHP
    (`App\\Domain\\Recipes\\Koszt\\SzacunekKosztuZCen`) — żeby ten skrypt
    nigdy nie trzymał osobnej kopii tej liczby (D-286, część 3: „mnożnik
    ma być jedną stałą z komentarzem i uzasadnieniem”, nie liczbą wpisaną
    po cichu w dwóch miejscach).
    """
    tresc = PLIK_MNOZNIKA.read_text(encoding="utf-8")
    dopasowanie = re.search(r"MNOZNIK_HURT_DETAL\s*=\s*([\d.]+)\s*;", tresc)

    if dopasowanie is None:
        raise SystemExit(
            f"Nie znalazłem stałej MNOZNIK_HURT_DETAL w {PLIK_MNOZNIKA} — sprawdź, czy ktoś nie zmienił jej nazwy.",
        )

    return float(dopasowanie.group(1))


def najnowszy_zasob() -> dict:
    with urllib.request.urlopen(API_ZASOBY, timeout=60) as odp:  # noqa: S310 — stały adres dane.gov.pl
        dane = json.load(odp)

    for wpis in dane.get("data", []):
        atrybuty = wpis["attributes"]
        tytul = str(atrybuty.get("title", ""))
        if tytul.lower().startswith("rynek owoc") and atrybuty.get("format") == "xlsx":
            return atrybuty

    raise SystemExit("Nie znalazłem najnowszego biuletynu „Rynek owoców i warzyw” (xlsx) w zbiorze 912.")


def pobierz_xlsx(url: str) -> bytes:
    with urllib.request.urlopen(url, timeout=120) as odp:  # noqa: S310 — adres wzięty z odpowiedzi API dane.gov.pl
        return odp.read()


def okres_z_tytulu(tytul: str) -> str:
    """„Rynek owoców i warzyw - notowania za okres: 14-22.09.2026 r.” → „notowania 14-22.09.2026”.

    Tytuł zasobu jest prostszy i stabilniejszy niż szukanie konkretnej
    komórki w arkuszu — układ arkuszy w tym biuletynie zmienia się
    (np. arkusz „ZAKUP WARZ DETAL - DO 2 KG” w ogóle nie ma własnej
    komórki z opisem tygodnia, w przeciwieństwie do „HURT WARZ”).
    """
    dopasowanie = re.search(r"notowania za okres:\s*([\d.\s-]+)\s*r\.?", tytul, re.IGNORECASE)
    if dopasowanie is None:
        raise SystemExit(f"Nie umiem odczytać okresu notowania z tytułu zasobu: {tytul!r}")

    return f"notowania {dopasowanie.group(1).strip()}"


def wyciagnij_ceny(bajty: bytes) -> dict[str, float]:
    """Zwraca {etykieta: cena_zl_za_kg} z arkusza „ZAKUP WARZ DETAL - DO 2 KG”."""
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover — komunikat dla człowieka, nie test
        raise SystemExit("Brakuje pakietu openpyxl. Zainstaluj: pip install openpyxl") from e

    skoroszyt = openpyxl.load_workbook(io.BytesIO(bajty), data_only=True)
    if ARKUSZ not in skoroszyt.sheetnames:
        raise SystemExit(f"Biuletyn nie ma arkusza „{ARKUSZ}” — sprawdź, czy MRiRW nie zmieniło układu pliku.")

    arkusz = skoroszyt[ARKUSZ]
    wiersze = list(arkusz.iter_rows(values_only=True))
    ceny: dict[str, float] = {}
    w_sekcji = False

    for wiersz in wiersze:
        pierwsza_niepusta = next((v for v in wiersz if v is not None), None)

        if pierwsza_niepusta == SEKCJA:
            w_sekcji = True
            continue

        if not w_sekcji:
            continue

        if pierwsza_niepusta is None:
            break  # pusty wiersz kończy sekcję

        etykieta = wiersz[1] if len(wiersz) > 1 else None
        cena_100kg = wiersz[2] if len(wiersz) > 2 else None

        if not isinstance(etykieta, str) or etykieta in ("TOWAR", ""):
            continue

        if isinstance(cena_100kg, (int, float)):
            ceny[etykieta.strip()] = round(float(cena_100kg) / 100, 2)

    return ceny


def wyciagnij_ceny_hurt(bajty: bytes) -> dict[str, float]:
    """Zwraca {etykieta: średnia_cena_hurtowa_zł} z arkusza „HURT WARZ”.

    Cena to ŚREDNIA z min–max KAŻDEGO z pięciu rynków (`RYNKI_HURT`), a
    potem średnia arytmetyczna tych pięciu wartości — dokładnie tak, jak
    przy ręcznym liczeniu opisanym w `docs/DECISIONS.md`, D-286 część 3
    („Uzupełnienie z 26.09.2026”). Układ kolumn: TOWAR, potem para
    (min, max) dla każdego z pięciu rynków, w kolejności `RYNKI_HURT`.

    Rynek bez notowania w danym tygodniu (pusta komórka, „-”, „nld”) jest
    POMIJANY przy liczeniu średniej — nie liczy się jako zero, bo zero
    zaniżałoby cenę pozostałych rynków, które faktycznie notowały.
    """
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover — komunikat dla człowieka, nie test
        raise SystemExit("Brakuje pakietu openpyxl. Zainstaluj: pip install openpyxl") from e

    skoroszyt = openpyxl.load_workbook(io.BytesIO(bajty), data_only=True)
    if ARKUSZ_HURT not in skoroszyt.sheetnames:
        raise SystemExit(f"Biuletyn nie ma arkusza „{ARKUSZ_HURT}” — sprawdź, czy MRiRW nie zmieniło układu pliku.")

    arkusz = skoroszyt[ARKUSZ_HURT]
    wiersze = list(arkusz.iter_rows(values_only=True))
    ceny: dict[str, float] = {}
    w_sekcji = False
    liczba_kolumn = 2 * len(RYNKI_HURT)

    for wiersz in wiersze:
        pierwsza_niepusta = next((v for v in wiersz if v is not None), None)

        if pierwsza_niepusta == SEKCJA_HURT:
            w_sekcji = True
            continue

        if not w_sekcji:
            continue

        if pierwsza_niepusta is None:
            break  # pusty wiersz kończy sekcję (dalej byłaby np. „WARZYWA z importu”)

        etykieta = wiersz[1] if len(wiersz) > 1 else None

        if not isinstance(etykieta, str) or etykieta in ("TOWAR", ""):
            continue

        wartosci = list(wiersz[2:2 + liczba_kolumn])
        srednie_rynkow = []

        for i in range(len(RYNKI_HURT)):
            minimum = wartosci[i * 2] if i * 2 < len(wartosci) else None
            maksimum = wartosci[i * 2 + 1] if i * 2 + 1 < len(wartosci) else None

            if isinstance(minimum, (int, float)) and isinstance(maksimum, (int, float)):
                srednie_rynkow.append((float(minimum) + float(maksimum)) / 2)

        if srednie_rynkow:
            ceny[etykieta.strip()] = sum(srednie_rynkow) / len(srednie_rynkow)

    return ceny


def _pl(liczba: float, miejsca: int) -> str:
    """`3.928` (2 miejsca) → `"3,93"` — separator dziesiętny jak w reszcie CSV."""
    return f"{liczba:.{miejsca}f}".replace(".", ",")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--zapisz", action="store_true", help="zapisz zmiany do pliku CSV")
    args = parser.parse_args()

    zasob = najnowszy_zasob()
    bajty = pobierz_xlsx(str(zasob["download_url"]))
    ceny = wyciagnij_ceny(bajty)
    ceny_hurt = wyciagnij_ceny_hurt(bajty)
    mnoznik = mnoznik_hurt_detal()
    okres = okres_z_tytulu(str(zasob.get("title", "")))
    rok = str(zasob.get("data_date", ""))[:4] or "brak-daty"

    with PLIK.open(newline="", encoding="utf-8") as f:
        czytnik = csv.DictReader(f)
        naglowek = czytnik.fieldnames or []
        wiersze = list(czytnik)

    brak = []
    for w in wiersze:
        if not w["zrodlo"].startswith("MRiRW"):
            continue

        klucz = w["klucz"]

        if klucz in MAPA:
            etykieta = MAPA[klucz]

            if etykieta not in ceny:
                brak.append(klucz)
                continue

            nowa = f"{ceny[etykieta]:.2f}"
            if nowa != w["cena_zl"] or w["okres"] != rok:
                print(f"{klucz}: {w['cena_zl']} zł ({w['okres']}) → {nowa} zł ({rok})")

            w["cena_zl"] = nowa
            w["okres"] = rok
            w["zrodlo"] = (
                f"MRiRW, ZSRIR: {okres}, średnia cena zakupu warzyw przez detal "
                f"(opak. do 2 kg), {etykieta.lower()} - za 1kg (z zł/100 kg)"
            )
            continue

        if klucz in MAPA_HURT:
            etykieta, jednostka = MAPA_HURT[klucz]

            if etykieta not in ceny_hurt:
                brak.append(klucz)
                continue

            srednia_hurt = ceny_hurt[etykieta]
            nowa = f"{round(srednia_hurt * mnoznik, 2):.2f}"
            precyzja_hurt = 3 if jednostka == "szt" else 2

            if nowa != w["cena_zl"] or w["okres"] != rok:
                print(f"{klucz}: {w['cena_zl']} zł ({w['okres']}) → {nowa} zł ({rok}) [hurt × {mnoznik}]")

            w["cena_zl"] = nowa
            w["okres"] = rok
            w["zrodlo"] = (
                f"MRiRW, ZSRIR: {okres}, szacunek z cen hurtowych "
                f'(arkusz "HURT WARZ", srednia min-max z 5 rynkow: {", ".join(RYNKI_HURT)}) '
                f"x {_pl(mnoznik, 1)} (przelicznik hurt-detal, D-286 cz. 3), "
                f"{etykieta.lower()} - hurt {_pl(srednia_hurt, precyzja_hurt)} zl/{jednostka}"
            )
            continue

        # Wiersz spoza obu map: warzywo, którego jeszcze nie obsługujemy
        # automatycznie — nowy klucz dopisuje się ręcznie razem z dopasowaniem
        # tekstu składnika (patrz docstring modułu).
        continue

    if brak:
        print(
            f"\nBiuletyn nie podał ceny dla: {', '.join(brak)} — te ceny zostają bez zmian.",
            file=sys.stderr,
        )

    if args.zapisz:
        with PLIK.open("w", newline="", encoding="utf-8") as f:
            pisarz = csv.DictWriter(f, fieldnames=naglowek, lineterminator="\n")
            pisarz.writeheader()
            pisarz.writerows(wiersze)
        print(f"\nZapisano {PLIK}. Przejrzyj różnice (git diff) i wczytaj po wdrożeniu: php artisan kuking:ceny-skladnikow")
    else:
        print("\nPodgląd — niczego nie zapisano. Dodaj --zapisz.")

    return 1 if brak else 0


if __name__ == "__main__":
    sys.exit(main())
